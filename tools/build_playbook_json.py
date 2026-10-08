"""Convert the HEN NOC Congestion Playbook (xlsx) into playbook.json for the dashboard.
Usage: python tools/build_playbook_json.py playbook/HEN_NOC_Congestion_Playbook.xlsx playbook.json
Optional: add a column headed "ERCOT Alias" to the Line Lookup sheet (any column) to give a
comma-separated list of exact ERCOT constraint names that should map to that row.
"""
import sys, json, re, datetime
import openpyxl

src, dst = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src, data_only=True)

def sheet(prefix):
    for ws in wb.worksheets:
        if prefix.lower() in ws.title.lower():
            return ws
    raise SystemExit(f"Sheet containing '{prefix}' not found")

def clean(v):
    if v is None: return ""
    s = str(v).strip()
    return "" if s.lower() in ("none", "nan", "—") else s

def num(v):
    try: return float(v)
    except (TypeError, ValueError): return None

def prio(v):
    s = re.sub(r"^[^A-Za-z]+", "", clean(v)).lower()
    for key, name in (("active", "Active"), ("watch", "Watch"), ("low", "Low"), ("new", "New")):
        if s.startswith(key):
            return name
    return None

# --- Shift factors per site (sheet: Shift Factors) ---
sf_ws = sheet("Shift Factors")
hdr = [clean(c.value) for c in sf_ws[3]]
site_cols = {}
for i, h in enumerate(hdr):
    if i >= 4 and h and not re.search(r"range|notable", h, re.I):
        site_cols[i] = h
sf = {}
for row in sf_ws.iter_rows(min_row=4, values_only=True):
    name = clean(row[0])
    if not name: continue
    sites = {site_cols[i]: num(row[i]) for i in site_cols if num(row[i]) is not None}
    sf[name.upper()] = {"sites": sites, "totalShadow": num(row[2])}

# --- Hour guide ---
hg_ws = sheet("Hour Guide")
hg_hdr = [clean(c.value) for c in hg_ws[3]]
hg = {}
for row in hg_ws.iter_rows(min_row=4, values_only=True):
    name = clean(row[0])
    if not name: continue
    hg[name.upper()] = [num(row[3 + h]) or 0 for h in range(24)]

# --- Line lookup ---
ll = sheet("Line Lookup")
head = [clean(c.value).replace("\n", " ") for c in ll[3]]
alias_col = next((i for i, h in enumerate(head) if "alias" in h.lower()), None)
lines = []
for row in ll.iter_rows(min_row=4, values_only=True):
    p = prio(row[1])
    name = clean(row[0])
    if not name or not p or "detailed" in clean(row[1]).lower():
        continue
    key = name.upper()
    aliases = [a.strip() for a in clean(row[alias_col]).split(",") if a.strip()] if alias_col is not None else []
    lines.append({
        "line": name,
        "priority": p,
        "season": clean(row[2]),
        "peakHE": clean(row[3]),
        "trigger": clean(row[4]),
        "primaryNode": clean(row[5]),
        "primarySF": num(row[6]),
        "mcc500": num(row[7]), "mcc1000": num(row[8]), "mcc2000": num(row[9]),
        "otherNodes": clean(row[10]),
        "holdStrategy": clean(row[11]),
        "notes": clean(row[12]),
        "aliases": aliases,
        "siteSF": sf.get(key, {}).get("sites", {}),
        "hourly": hg.get(key),
    })

# --- Disclaimer text ---
disc = [clean(r[1]) for r in sheet("Disclaimer").iter_rows(values_only=True) if clean(r[1])]

out = {
    "generated": datetime.datetime.utcnow().isoformat() + "Z",
    "source": src.split("/")[-1],
    "disclaimer": disc,
    "lines": lines,
}
json.dump(out, open(dst, "w"), indent=1)
print(f"Wrote {len(lines)} lines to {dst}")
