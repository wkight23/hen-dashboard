"""One-time crosswalk builder: ERCOT constraint names -> HEN playbook lines.

Pulls the last N days of binding constraints (NP6-86-CD), then scores every ERCOT
constraint against every line in playbook.json and writes crosswalk.xlsx with:
  Crosswalk              one block per playbook line with its best ERCOT candidates
  Aliases_to_paste       High-confidence pairs in the exact format of the Aliases sheet
  Frequent_not_in_playbook   busiest ERCOT constraints that match nothing in the playbook
Env: ERCOT_USERNAME, ERCOT_PASSWORD, ERCOT_SUBSCRIPTION_KEY
Usage: python tools/build_crosswalk.py --days 180 [--playbook playbook.json] [--out crosswalk.xlsx]
"""
import os, re, sys, json, time, argparse, datetime
from collections import defaultdict

# ------------------------------------------------------------------ scoring
def norm(s): return re.sub(r"[^A-Z0-9]", "", str(s or "").upper())
def alpha(s): return re.sub(r"\d+$", "", norm(s))          # drop trailing digits (CATARI1 -> CATARI)
def pr(a, b):
    # short station codes (STP, WAP, ZEN...) must match exactly; longer ones may match by prefix
    if not a or not b: return False
    if min(len(a), len(b)) < 4: return a == b
    return a.startswith(b) or b.startswith(a)

def parse_line(name):
    m = re.search(r"(\d+)\s*kv", name, re.I)
    kv = int(m.group(1)) if m else None
    base = re.sub(r"\s*\d+\s*kv\s*$", "", name, flags=re.I).strip()
    parts = [norm(p) for p in base.split("-") if norm(p)]
    return parts, kv, norm(base)

def kv_adjust(line_kv, e_kv):
    try: e_kv = int(float(e_kv))
    except (TypeError, ValueError): return 0
    if line_kv is None: return 0
    return 5 if e_kv == line_kv else -15

def score(parts, lkv, lbase, e):
    nm, fr, to = norm(e["name"]), norm(e["from"]), norm(e["to"])
    toks = [alpha(t) for t in re.split(r"[^A-Za-z0-9]+", str(e["name"])) if alpha(t)]
    s = 0
    if len(parts) >= 2:
        p, q = parts[0], parts[1]
        if nm == lbase: s = 100
        if fr and to:
            if pr(fr, p) and pr(to, q): s = max(s, 90)
            elif pr(fr, q) and pr(to, p): s = max(s, 70)
        if len(toks) >= 2:
            if pr(toks[0], p) and pr(toks[1], q): s = max(s, 80)
            elif pr(toks[0], q) and pr(toks[1], p): s = max(s, 60)
    elif len(parts) == 1:
        p = parts[0]
        if nm == p: s = 95
        elif toks and pr(toks[0], p): s = 70
        elif pr(fr, p) or pr(to, p): s = 55
        elif any(pr(t, p) for t in toks): s = 60
    if s == 0: return 0
    return max(0, min(100, s + kv_adjust(lkv, e["kv"])))

def label(s): return "High" if s >= 85 else "Medium" if s >= 60 else "Low"

def match_all(lines, agg):
    ercot = list(agg.values())
    out, best_for_ercot = {}, defaultdict(int)
    for l in lines:
        parts, lkv, lbase = parse_line(l["line"])
        cands = []
        for e in ercot:
            sc = score(parts, lkv, lbase, e)
            if sc >= 45:
                cands.append((sc, e))
                best_for_ercot[id(e)] = max(best_for_ercot[id(e)], sc)
        cands.sort(key=lambda x: (-x[0], -x[1]["intervals"]))
        out[l["line"]] = cands[:5]
    return out, best_for_ercot

# ------------------------------------------------------------------ workbook
def write_workbook(lines, agg, path, days):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    cand, best = match_all(lines, agg)
    wb = openpyxl.Workbook()
    hdr_fill = PatternFill("solid", fgColor="1F3A5F"); hdr_font = Font(bold=True, color="FFFFFF")
    fills = {"High": "C6EFCE", "Medium": "FFEB9C", "Low": "F8CBAD", "None": "D9D9D9"}

    def header(ws, cols):
        ws.append(cols)
        for c in ws[1]: c.fill, c.font = hdr_fill, hdr_font
        ws.freeze_panes = "A2"

    ws = wb.active; ws.title = "Crosswalk"
    header(ws, ["Playbook Line", "Priority", "Rank", "Confidence", "Score", "ERCOT Constraint Name",
                "From", "To", "kV", "Intervals Bound", "Days Seen", "Total Shadow $", "Max Shadow $", "Contingencies Seen"])
    for l in lines:
        cs = cand[l["line"]]
        if not cs:
            ws.append([l["line"], l["priority"], "", "None", "", "NO CANDIDATE FOUND IN LAST %d DAYS" % days])
            for c in ws[ws.max_row]: c.fill = PatternFill("solid", fgColor=fills["None"])
            continue
        for i, (sc, e) in enumerate(cs, 1):
            ws.append([l["line"] if i == 1 else "", l["priority"] if i == 1 else "", i, label(sc), sc, e["name"],
                       e["from"], e["to"], e["kv"], e["intervals"], len(e["days"]),
                       round(e["sum_sp"], 2), round(e["max_sp"], 2), ", ".join(sorted(e["conts"])[:4])])
            ws.cell(ws.max_row, 4).fill = PatternFill("solid", fgColor=fills[label(sc)])
            if i == 1:
                ws.cell(ws.max_row, 1).font = Font(bold=True)
    for col, w in zip("ABCDEFGHIJKLMN", [34, 10, 6, 12, 7, 26, 12, 12, 7, 14, 10, 16, 14, 40]):
        ws.column_dimensions[col].width = w

    ws2 = wb.create_sheet("Aliases_to_paste")
    header(ws2, ["ERCOT Constraint Name", "Playbook Line"])
    seen = set()
    for l in lines:
        for sc, e in cand[l["line"]]:
            if sc >= 85 and (e["name"], l["line"]) not in seen:
                seen.add((e["name"], l["line"])); ws2.append([e["name"], l["line"]])
    ws2.column_dimensions["A"].width = 30; ws2.column_dimensions["B"].width = 36

    ws3 = wb.create_sheet("Frequent_not_in_playbook")
    header(ws3, ["ERCOT Constraint Name", "From", "To", "kV", "Intervals Bound", "Days Seen",
                 "Total Shadow $", "Max Shadow $", "Contingencies Seen"])
    unm = [e for e in agg.values() if best.get(id(e), 0) < 60]
    unm.sort(key=lambda e: -e["sum_sp"])
    for e in unm[:100]:
        ws3.append([e["name"], e["from"], e["to"], e["kv"], e["intervals"], len(e["days"]),
                    round(e["sum_sp"], 2), round(e["max_sp"], 2), ", ".join(sorted(e["conts"])[:4])])
    for col, w in zip("ABCDEFGHI", [30, 12, 12, 7, 14, 10, 16, 14, 40]):
        ws3.column_dimensions[col].width = w
    wb.save(path)
    hi = sum(1 for l in lines if cand[l["line"]] and cand[l["line"]][0][0] >= 85)
    md = sum(1 for l in lines if cand[l["line"]] and 60 <= cand[l["line"]][0][0] < 85)
    no = sum(1 for l in lines if not cand[l["line"]] or cand[l["line"]][0][0] < 60)
    print(f"{len(lines)} playbook lines: {hi} high-confidence, {md} medium, {no} low/none. Wrote {path}")

# ------------------------------------------------------------------ ERCOT pull
TOKEN_URL = "https://ercotb2c.b2clogin.com/ercotb2c.onmicrosoft.com/B2C_1_PUBAPI-ROPC-FLOW/oauth2/v2.0/token"
CLIENT_ID = "fec253ea-0d06-4272-a5e6-b478baeecd70"
API = "https://api.ercot.com/api/public-reports/np6-86-cd/shdw_prices_bnd_trns_const"

def pull(days):
    import requests
    key = os.environ["ERCOT_SUBSCRIPTION_KEY"]
    state = {"tok": None, "t": 0, "size": 50000}

    def token():
        if state["tok"] and time.time() - state["t"] < 25 * 60: return state["tok"]
        r = requests.post(TOKEN_URL, data={
            "username": os.environ["ERCOT_USERNAME"], "password": os.environ["ERCOT_PASSWORD"],
            "grant_type": "password", "scope": f"openid {CLIENT_ID} offline_access",
            "client_id": CLIENT_ID, "response_type": "id_token"}, timeout=30)
        r.raise_for_status(); state["tok"], state["t"] = r.json()["id_token"], time.time()
        return state["tok"]

    def get(params):
        for attempt in range(8):
            time.sleep(2.2)                                   # stay under ERCOT's ~30 requests/min
            r = requests.get(API, params=params, timeout=120,
                             headers={"Authorization": f"Bearer {token()}", "Ocp-Apim-Subscription-Key": key})
            if r.status_code == 200: return r.json()
            if r.status_code == 400 and state["size"] > 5000:
                state["size"] = 5000; params["size"] = 5000; print("page size reduced to 5000"); continue
            if r.status_code == 401: state["tok"] = None
            print(f"  HTTP {r.status_code}; retry {attempt+1}"); time.sleep(5 * (attempt + 1))
        raise RuntimeError("giving up on request")

    def g(d, *keys):
        low = {k.lower(): v for k, v in d.items()}
        for k in keys:
            if k.lower() in low: return low[k.lower()]

    agg, failed = {}, []
    try:
        from zoneinfo import ZoneInfo
        today = datetime.datetime.now(ZoneInfo("America/Chicago")).date()
    except Exception:
        today = (datetime.datetime.utcnow() - datetime.timedelta(hours=5)).date()
    for n in range(days, -1, -1):
        d = today - datetime.timedelta(days=n)
        try:
            page, rows = 1, 0
            while True:
                j = get({"SCEDTimestampFrom": d.strftime("%Y-%m-%dT00:00:00"), "SCEDTimestampTo": d.strftime("%Y-%m-%dT23:59:59"),
                         "size": state["size"], "page": page})
                names = [f["name"] for f in j.get("fields", [])]
                for row in j.get("data", []):
                    r = dict(zip(names, row)) if isinstance(row, list) else row
                    name = g(r, "constraintName") or ""; fr = g(r, "fromStation") or ""; to = g(r, "toStation") or ""
                    kv = g(r, "fromStationkV", "kV")
                    k = (name, fr, to, kv)
                    e = agg.setdefault(k, {"name": name, "from": fr, "to": to, "kv": kv, "intervals": 0,
                                           "sum_sp": 0.0, "max_sp": 0.0, "days": set(), "conts": set()})
                    sp = float(g(r, "shadowPrice") or 0)
                    e["intervals"] += 1; e["sum_sp"] += sp; e["max_sp"] = max(e["max_sp"], sp)
                    e["days"].add(str(d)); c = g(r, "contingencyName")
                    if c and len(e["conts"]) < 12: e["conts"].add(str(c))
                    rows += 1
                if page >= int(j.get("_meta", {}).get("totalPages", 1)): break
                page += 1
            print(f"{d}: {rows} rows, {len(agg)} distinct constraints so far", flush=True)
        except Exception as ex:
            print(f"{d}: FAILED ({ex})"); failed.append(str(d))
    if failed: print("Days that failed:", ", ".join(failed))
    return agg

if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--days", type=int, default=180)
    ap.add_argument("--playbook", default="playbook.json")
    ap.add_argument("--out", default="crosswalk.xlsx")
    ap.add_argument("--from-agg", help="re-score from a saved agg.json instead of calling ERCOT")
    a = ap.parse_args()
    lines = json.load(open(a.playbook))["lines"]
    if a.from_agg:
        agg = {i: {**e, "days": set(e["days"]), "conts": set(e["conts"])} for i, e in enumerate(json.load(open(a.from_agg)))}
    else:
        agg = pull(a.days)
        json.dump([{**e, "days": sorted(e["days"]), "conts": sorted(e["conts"])} for e in agg.values()], open("agg.json", "w"))
    write_workbook(lines, agg, a.out, a.days)
